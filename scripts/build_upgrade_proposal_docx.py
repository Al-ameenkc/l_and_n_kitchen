"""Build L&N Kitchen menu website upgrade proposal as a styled DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "proposal-screenshots"
OUT = ROOT / "L_and_N_Menu_Website_Upgrade_Proposal.docx"

# Brand palette
BG = "0A0A0B"
CARD = "141416"
GREEN = "16A34A"
GREEN_SOFT = "4ADE80"
WHITE = "FAFAFA"
MUTED = "A1A1AA"
LINE = "27272A"
DARK_TEXT = "18181B"
BODY = "3F3F46"


def rgb(hex6: str) -> RGBColor:
    h = hex6.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_cell_shading(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_paragraph_shading(paragraph, hex_color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    p_pr.append(shading)


def add_run(paragraph, text, *, size=11, bold=False, color=DARK_TEXT, font="Calibri"):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return run


def add_heading_block(doc, title: str, subtitle: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, title, size=22, bold=True, color=WHITE)
    if subtitle:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(12)
        add_run(p2, subtitle, size=11, color=MUTED)
    # green accent bar row
    accent = doc.add_table(rows=1, cols=1)
    set_cell_shading(accent.rows[0].cells[0], GREEN)
    accent.rows[0].cells[0].height = Cm(0.15)
    doc.add_paragraph()


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        add_run(p, item, size=11, color=BODY)


def add_image_row(doc, images: list[tuple[str, str]], width_in: float = 2.1) -> None:
    if not images:
        return
    table = doc.add_table(rows=2, cols=len(images))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, (path, caption) in enumerate(images):
        cell = table.rows[0].cells[col]
        set_cell_shading(cell, CARD)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path = SHOTS / path
        if img_path.exists():
            run = p.add_run()
            run.add_picture(str(img_path), width=Inches(width_in))
        else:
            add_run(p, f"[{path}]", size=9, color=MUTED)
        cap = table.rows[1].cells[col]
        cp = cap.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(cp, caption, size=9, bold=True, color=MUTED)
    doc.add_paragraph()


def add_card_table(doc, cards: list[tuple[str, list[str]]], accent_first=False) -> None:
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (title, lines) in enumerate(cards):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, CARD if not (accent_first and i == 0) else "1A2E1F")
        p = cell.paragraphs[0]
        add_run(p, title, size=12, bold=True, color=GREEN_SOFT if accent_first and i == 0 else WHITE)
        for line in lines:
            lp = cell.add_paragraph()
            lp.paragraph_format.space_after = Pt(2)
            add_run(lp, line, size=10, color=MUTED)
    doc.add_paragraph()


def add_cost_table(doc) -> None:
    rows = [
        ("UX / UI design & brand landing page", "₦300,000", "16%"),
        ("Public website + desktop menu browse", "₦350,000", "19%"),
        ("Guest journey upgrades (shared order, status, filters)", "₦400,000", "22%"),
        ("Waiter & kitchen ops (KDS, alerts, roles, reports)", "₦450,000", "24%"),
        ("Analytics, staff training & 30-day support", "₦200,000", "11%"),
        ("Project contingency & QA buffer", "₦150,000", "8%"),
    ]
    table = doc.add_table(rows=1 + len(rows) + 1, cols=3)
    table.style = "Table Grid"
    headers = ["Workstream", "Amount (NGN)", "Share"]
    for c, h in enumerate(headers):
        cell = table.rows[0].cells[c]
        set_cell_shading(cell, BG)
        add_run(cell.paragraphs[0], h, size=10, bold=True, color=GREEN_SOFT)
    for r, (a, b, c) in enumerate(rows, start=1):
        for col, text in enumerate([a, b, c]):
            cell = table.rows[r].cells[col]
            set_cell_shading(cell, "F4F4F5" if r % 2 else WHITE)
            align = WD_ALIGN_PARAGRAPH.LEFT
            if col == 1:
                align = WD_ALIGN_PARAGRAPH.RIGHT
            elif col == 2:
                align = WD_ALIGN_PARAGRAPH.CENTER
            p = cell.paragraphs[0]
            p.alignment = align
            add_run(p, text, size=10, bold=col == 0, color=DARK_TEXT if col == 0 else BODY)
    total = table.rows[-1]
    for cell in total.cells:
        set_cell_shading(cell, GREEN)
    add_run(total.cells[0].paragraphs[0], "TOTAL PACKAGE", size=12, bold=True, color=WHITE)
    tp = total.cells[1].paragraphs[0]
    tp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(tp, "₦1,850,000", size=16, bold=True, color=WHITE)
    total.cells[2].paragraphs[0].clear()
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # Cover
    cover = doc.add_table(rows=1, cols=1)
    cell = cover.rows[0].cells[0]
    set_cell_shading(cell, BG)
    cell.height = Cm(8)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(48)
    add_run(p, "L&N KITCHEN", size=14, bold=True, color=GREEN_SOFT)
    t = cell.add_paragraph()
    add_run(t, "Menu Website Upgrade Proposal", size=28, bold=True, color=WHITE)
    s = cell.add_paragraph()
    s.paragraph_format.space_before = Pt(12)
    add_run(
        s,
        "From swipeable digital menu to a full guest-to-kitchen ordering experience.\n"
        "An Oasis of Pleasure — modernised for speed, control, and hospitality.",
        size=12,
        color=MUTED,
    )
    m = cell.add_paragraph()
    m.paragraph_format.space_before = Pt(24)
    add_run(m, "Fixed investment: ₦1,850,000  ·  August 2026", size=11, color=MUTED)
    bar = doc.add_table(rows=1, cols=1)
    set_cell_shading(bar.rows[0].cells[0], GREEN)
    doc.add_page_break()

    add_heading_block(doc, "Executive summary")
    add_bullets(
        doc,
        [
            "L&N Kitchen already runs a mobile-first swipe menu with live orders, waiter assign, and admin control.",
            "This proposal upgrades the system into a branded public website and deeper operations tooling.",
            "All phases below are included in a fixed package of ₦1,850,000.",
            "Screenshots in this document were captured from the live Grill category menu.",
        ],
    )

    add_heading_block(doc, "Customer experience (mobile)", "Grill category · swipe deck, dish detail, wish list")
    add_image_row(
        doc,
        [
            ("01-customer-menu-home.png", "Swipe menu — Grill"),
            ("03-customer-dish-detail.png", "Dish detail"),
            ("04-customer-wishlist-sheet.png", "Wish list"),
        ],
        width_in=1.95,
    )

    add_heading_block(doc, "Admin dashboard", "Login, live orders, menu CMS, and waiters")
    add_image_row(
        doc,
        [
            ("05-admin-login.png", "Admin login"),
            ("06-admin-orders.png", "Live orders"),
        ],
        width_in=3.0,
    )
    dishes_img = "08b-admin-dishes-list.png" if (SHOTS / "08b-admin-dishes-list.png").exists() else "08-admin-dishes.png"
    add_image_row(
        doc,
        [
            ("07-admin-categories.png", "Categories"),
            (dishes_img, "Dishes"),
            ("09-admin-waiters.png", "Waiters"),
        ],
        width_in=1.95,
    )

    add_heading_block(doc, "Current state", "Already live in the L&N Kitchen system")
    add_card_table(
        doc,
        [
            (
                "Guest menu",
                [
                    "Mobile swipe card deck",
                    "Category carousel & search",
                    "Wish list + trash restore",
                    "Dish detail & allergens",
                    "QR share for waiter pickup",
                    "Live sync after assign",
                    "Session reset after payment",
                ],
            ),
            (
                "Floor operations",
                [
                    "Waiter scans guest QR",
                    "Table / lounge capture",
                    "Assign to admin (one-shot)",
                    "Order image download",
                    "Staff ID waiter profiles",
                ],
            ),
            (
                "Admin control",
                [
                    "Orders dashboard",
                    "Accept / mark paid",
                    "Edit live order items",
                    "Categories & dishes CMS",
                    "Availability toggles",
                    "Waiter management",
                ],
            ),
        ],
        accent_first=True,
    )

    add_heading_block(doc, "Opportunity & vision")
    add_bullets(
        doc,
        [
            "Turn the in-venue phone menu into a branded L&N web destination.",
            "Reduce paper and chat-order mistakes with one source of truth.",
            "Give kitchen and lounge shared live status from guest to paid.",
            "Use order data for bestsellers, peak hours, and staff insight.",
            "North star: digital hospitality that feels as curated as the food.",
        ],
    )

    add_heading_block(doc, "Upgrade pillars")
    add_card_table(
        doc,
        [
            ("01  Brand & web", ["Public landing + menu site", "SEO / share previews", "Desktop-friendly browse", "Story & atmosphere"]),
            ("02  Guest journey", ["Table join via QR", "Multi-guest shared order", "Status timeline", "Receipts & feedback"]),
            ("03  Ops excellence", ["Kitchen display (KDS)", "Printer / alerts", "Shift reports", "Role permissions"]),
            ("04  Growth", ["Promos & set menus", "VIP / lounge packages", "Analytics dashboard", "Loyalty hooks"]),
        ],
        accent_first=True,
    )

    add_heading_block(doc, "Phased roadmap", "Everything below is covered by the ₦1,850,000 package")
    phases = [
        ("Phase 1 — Polish (Weeks 1–3)", ["Public brand landing", "Desktop menu browse", "Status UX clarity", "Analytics basics"]),
        ("Phase 2 — Shared dining (Weeks 4–7)", ["Multi-guest table order", "Guest status timeline", "Waiter alerts", "Dietary filters"]),
        ("Phase 3 — Kitchen (Weeks 8–12)", ["Kitchen display", "Print / notify hooks", "Roles & audit", "Shift reports"]),
        ("Phase 4 — Growth (Weeks 12–14)", ["Promos & set menus", "VIP packages", "Loyalty hooks", "SEO polish"]),
    ]
    add_card_table(doc, phases)

    add_heading_block(
        doc,
        "Cost breakdown — ₦1,850,000",
        "Fixed package investment (One Million, Eight Hundred and Fifty Thousand Naira)",
    )
    add_cost_table(doc)
    note = doc.add_paragraph()
    add_run(
        note,
        "Includes delivery of Phases 1–4. Hosting and domain billed separately if not already covered.",
        size=10,
        color=MUTED,
    )

    add_heading_block(doc, "Suggested payment schedule")
    add_card_table(
        doc,
        [
            ("On kickoff · 40%", ["₦740,000", "Brand assets collected, Phase 1 starts"]),
            ("Phase 1–2 demo · 30%", ["₦555,000", "Public site + shared dining live"]),
            ("Phase 3–4 handoff · 30%", ["₦555,000", "Kitchen ops + growth features + training"]),
        ],
        accent_first=True,
    )

    add_heading_block(doc, "Expected business outcomes")
    add_card_table(
        doc,
        [
            ("Faster turns", ["Clearer orders and paid resets reduce dead time between guests"]),
            ("Fewer errors", ["One digital source of truth vs handwritten or chat tickets"]),
            ("Higher attach", ["Upsell-ready UI for VIP, sides, and set experiences"]),
            ("Brand lift", ["A website that feels like L&N — not a generic PDF menu"]),
            ("Staff clarity", ["Waiters and admin share the same live order state"]),
            ("Insight", ["Know what sells, when, and which tables convert wish lists"]),
        ],
    )

    add_heading_block(doc, "Recommended next steps")
    add_bullets(
        doc,
        [
            "Approve the ₦1,850,000 package and preferred payment schedule.",
            "Run a brand workshop — logo, photos, hours, location, VIP / lounge story.",
            "Kick off Phase 1 — public site and desktop browse as the first visible upgrade.",
            "Pilot on the floor — train waiters and capture friction notes for Phases 2–3.",
        ],
    )

    # Close
    doc.add_page_break()
    close = doc.add_table(rows=1, cols=1)
    cell = close.rows[0].cells[0]
    set_cell_shading(cell, BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(36)
    add_run(p, "L&N KITCHEN", size=12, bold=True, color=GREEN_SOFT)
    t = cell.add_paragraph()
    add_run(t, "Let’s make the menu feel like the oasis guests already love.", size=20, bold=True, color=WHITE)
    f = cell.add_paragraph()
    f.paragraph_format.space_before = Pt(16)
    add_run(
        f,
        "Fixed investment: ₦1,850,000  ·  Phases 1–4 included  ·  Ready to kick off on approval",
        size=11,
        color=MUTED,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
